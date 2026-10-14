from pathlib import Path
from datetime import datetime, date
import csv, re, json, hashlib, zipfile, codecs
import openpyxl
from .config import SCHEMA, SCHEMAS, RUNTIME, INDUSTRIES
from .db import db, uid, now, dumps
from .storage import store

MISSING={'','报告未明确','报告未区分','未区分','未明确','不详','未知','暂无','无数据','--','N/A','nan'}
def missing(v): return v is None or str(v).strip() in MISSING
def num(v):
    if missing(v): return None
    if isinstance(v,(float,int)): return float(v)
    text=str(v).replace(',','').replace('，','').strip()
    m=re.fullmatch(r'([+-]?\d+(?:\.\d+)?)\s*(?:人|起|家|次|条|%|万元|亿元|元|吨)?',text)
    return float(m[1]) if m else None
def money(v):
    n=num(v)
    if n is None: return None
    t=str(v)
    if '亿元' in t: return n*10000
    if '元' in t and '万元' not in t and '亿元' not in t: return n/10000
    return n
def normalized_date(v):
    if missing(v): return ''
    if isinstance(v,(datetime,date)): return v.strftime('%Y-%m-%d')
    t=str(v).strip().replace('年','-').replace('月','-').replace('日','').replace('/','-').replace('.','-')
    m=re.match(r'^(\d{4})-(\d{1,2})(?:-(\d{1,2}))?',t)
    if not m: return ''
    try: return date(int(m[1]),int(m[2]),int(m[3] or 1)).isoformat()
    except ValueError: return ''
def derived(dataset,data):
    industry=str(data.get('行业') or '').replace('工贸行业','工贸')
    region='/'.join(str(data.get(f) or '') for f in ['省份','地市','区县'])
    d=next((data.get(f) for f in ['发生时间','统计日期','统计期'] if data.get(f)),None)
    return industry,region,normalized_date(d)
def clean(v):
    if isinstance(v,(date,datetime)): return v.isoformat(sep=' ') if isinstance(v,datetime) else v.isoformat()
    return str(v).strip() if v is not None else None

def validate_record(dataset,data):
    s=SCHEMA[dataset]; extra=set(data)-set(s['fields'])
    if extra: raise ValueError('不属于原字段体系：'+','.join(sorted(extra)))
    if not any(not missing(v) for v in data.values()): raise ValueError('空记录')
    data={k:clean(v) for k,v in data.items()}
    if missing(data.get(s['fields'][0])): raise ValueError('缺少'+s['fields'][0])
    if dataset in ['accidents','enterprises']:
        i=str(data.get('行业') or '').replace('工贸行业','工贸')
        if i not in INDUSTRIES: raise ValueError('行业需属于六行业')
    for f,v in data.items():
        if v and len(v)>100000: raise ValueError(f+'超过字段长度限制')
        if v and len(v)>1900 and not any(k in f for k in ['原因','经过','整改','救援','备注','说明','来源链接']): raise ValueError(f+'超过字段长度限制')
        if f in ['死亡人数','受伤人数','失踪人数','重伤人数','轻伤人数'] and not missing(v):
            n=num(v)
            if n is None or n<0 or n!=int(n): raise ValueError(f+'需为非负整数或明确缺失标记')
        if f in ['发生时间','统计日期','统计期'] and not missing(v) and not normalized_date(v):
            if not (f=='统计期' and re.fullmatch(r'\d{4}(?:年|Q[1-4]|年第[一二三四]季度)',str(v))): raise ValueError(f+'日期格式无效')
    return data

def safe_zip(path):
    with zipfile.ZipFile(path) as z:
        total=sum(i.file_size for i in z.infolist())
        if total>2*1024**3 or any(i.file_size>100*1024**2 and i.file_size/max(i.compress_size,1)>300 for i in z.infolist()):
            raise ValueError('压缩包展开体积或压缩率超限')

def excel_rows(path):
    safe_zip(path)
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        for sheet in w:
            rows=sheet.iter_rows(values_only=True)
            top=[]
            for _ in range(3): top.append(next(rows,()))
            first=[str(v).strip() if v is not None else '' for v in top[0]]
            s=next((s for s in SCHEMAS if sheet.title in [s['sheet'],s['id'],s['label']]),None)
            if s is None:
                s=next((s for s in SCHEMAS if first and first[0]==s['fields'][0] and sum(x in s['fields'] for x in first)>3),None)
            if s is None: raise ValueError('无法匹配Sheet：'+sheet.title)
            if sum(v in s['fields'] for v in first if v)>len([v for v in first if v])*0.8:
                fields=first; data_rows=iter(top[1:]+list())
                def chain():
                    yield from data_rows
                    yield from rows
                source=chain(); start_row=2
            else:
                expected=s['fields']
                header=[next((str(top[r][c]).strip() for r in [2,1,0] if c<len(top[r]) and top[r][c] is not None),'') for c in range(len(expected))] if s['id']=='accidents' else [str(v).strip() for v in top[1]]
                if header!=expected: raise ValueError('多行表头与原字段标准不一致：'+sheet.title)
                fields=expected; source=rows; start_row=4
            for lineno,row in enumerate(source,start_row):
                if not any(v is not None and str(v).strip() for v in row): continue
                if any(not f and i<len(row) and row[i] is not None for i,f in enumerate(fields)): raise ValueError('存在未命名列')
                data={f:clean(row[i]) if i<len(row) else None for i,f in enumerate(fields) if f}
                yield s['id'],lineno,data
    finally: w.close()

def parse_to_stage(path,name,dataset=None):
    id=uid(); counts={}; previews={}; rejected=0; valid=0
    stage=RUNTIME/'staging'/f'{id}.jsonl'; errors=RUNTIME/'staging'/f'{id}.errors.jsonl'
    ext=Path(name).suffix.lower()
    if ext=='.xlsx': source=excel_rows(path)
    elif ext=='.csv':
        if dataset not in SCHEMA: raise ValueError('CSV需指定数据集')
        def csv_rows():
            with open(path,'rb') as f: sample=f.read(65536)
            encoding='utf-8-sig'
            try: codecs.getincrementaldecoder(encoding)().decode(sample,final=False)
            except UnicodeDecodeError: encoding='gb18030'
            with open(path,'r',encoding=encoding,newline='') as f:
                reader=csv.DictReader(f)
                fields=reader.fieldnames or []
                if len(set(fields))!=len(fields): raise ValueError('CSV表头存在重复字段')
                if set(fields)-set(SCHEMA[dataset]['fields']): raise ValueError('CSV表头包含原字段体系以外的字段')
                for line,row in enumerate(reader,2):
                    if None in row: raise ValueError('CSV列数与表头不一致')
                    yield dataset,line,row
        source=csv_rows()
    else: return {'id':id,'raw':True,'valid':0,'rejected':0,'counts':{},'previews':{}}
    with stage.open('w',encoding='utf-8') as out,errors.open('w',encoding='utf-8') as err:
        for ds,line,row in source:
            try:
                cleaned=validate_record(ds,row)
                out.write(dumps({'dataset':ds,'row':line,'data':cleaned})+'\n')
                valid+=1; counts[ds]=counts.get(ds,0)+1
                if len(previews.setdefault(ds,[]))<5: previews[ds].append(cleaned)
            except ValueError as e:
                rejected+=1; err.write(dumps({'dataset':ds,'row':line,'error':str(e),'data':row})+'\n')
    return {'id':id,'raw':False,'valid':valid,'rejected':rejected,'counts':counts,'previews':previews}

def extract_raw(path,name):
    ext=Path(name).suffix.lower()
    if ext=='.pdf':
        from pypdf import PdfReader
        r=PdfReader(path); return '\n'.join((p.extract_text() or '') for p in r.pages[:500])[:500000]
    if ext=='.docx':
        safe_zip(path)
        from docx import Document
        doc=Document(path); return '\n'.join(p.text for p in doc.paragraphs)[:500000]
    if ext in ['.html','.htm']:
        from bs4 import BeautifulSoup
        soup=BeautifulSoup(path.read_bytes()[:5*1024**2],'html.parser')
        for e in soup(['script','style']): e.decompose()
        return soup.get_text('\n',strip=True)[:500000]
    if ext in ['.txt','.csv']: return path.read_text(encoding='utf-8-sig',errors='replace')[:500000]
    return ''

def commit_import(job_id,stage_id,actor):
    from .jobs import update
    stage=db.get_doc('stage',stage_id)
    if not stage or stage.get('committed'): raise ValueError('导入批次不存在或已经提交')
    file_id=stage['file_id']; file=db.query('SELECT * FROM RISK_FILES WHERE ID=?',[file_id])[0]
    imported=0; duplicate=0; processed=0
    path=RUNTIME/'staging'/f'{stage_id}.jsonl'
    if path.exists():
        with path.open(encoding='utf-8') as f:
            while True:
                batch=[]
                for _ in range(500):
                    line=f.readline()
                    if not line: break
                    batch.append(json.loads(line))
                if not batch: break
                with db.transaction() as c:
                    for row in batch:
                        if db.add_record(row['dataset'],row['data'],file_id,c): imported+=1
                        else: duplicate+=1
                processed+=len(batch)
                update(job_id,stage='导入结构化数据',progress=min(90,10+int(processed/max(stage['valid'],1)*80)))
    db.execute('UPDATE RISK_FILES SET STATUS=? WHERE ID=?',['imported',file_id])
    stage.update(committed=True,imported=imported,duplicate=duplicate); db.put_doc('stage',stage_id,stage)
    db.audit(actor,'import',file_id,{'imported':imported,'duplicate':duplicate,'rejected':stage['rejected']})
    path.unlink(missing_ok=True)
    return {'file_id':file_id,'imported':imported,'duplicate':duplicate,'rejected':stage['rejected']}
